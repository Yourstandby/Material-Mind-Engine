"""生成 MME 周报 PPT（16:9，与 Web 周报同一套设计语言）。"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(r"D:\Pi\MME\MME-周报-2026-W42.pptx")

PAPER = RGBColor(0xF7, 0xF6, 0xF3)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x14, 0x13, 0x0F)
INK2 = RGBColor(0x3D, 0x3B, 0x35)
MUTED = RGBColor(0x8A, 0x85, 0x79)
RULE = RGBColor(0xE4, 0xE1, 0xD8)
RULE_STRONG = RGBColor(0xD3, 0xCF, 0xC3)
ACCENT = RGBColor(0xC2, 0x41, 0x0C)
OK = RGBColor(0x15, 0x80, 0x3D)

YAHEI = "Microsoft YaHei"
MONO = "Consolas"

W, H = 13.333, 7.5


def _font(run, size, color, bold=False, name=YAHEI, spc=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", name)
    if spc:
        rpr.set("spc", str(spc))


def text(slide, x, y, w, h, s, size=13, color=INK2, bold=False, name=YAHEI,
         align=PP_ALIGN.LEFT, spacing=1.25, anchor=MSO_ANCHOR.TOP, spc=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    run = p.add_run()
    run.text = s
    _font(run, size, color, bold, name, spc)
    return box


def block(slide, x, y, w, items, gap=0.16):
    """items: [(标题, 说明)]"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.4))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for title, desc in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = 1.2
        p.space_after = Pt(gap * 28)
        r1 = p.add_run()
        r1.text = title
        _font(r1, 13.5, INK, True)
        if desc:
            r2 = p.add_run()
            r2.text = "\n" + desc
            _font(r2, 11.5, MUTED)
    return box


def line(slide, x, y, w, color=RULE, thickness=0.75):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Pt(thickness))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def vline(slide, x, y, h, color=RULE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Pt(0.75), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAPER
    return slide


def header(slide, eyebrow, title):
    text(slide, 0.9, 0.62, 11, 0.25, eyebrow, 9.5, MUTED, name=MONO, spc=180)
    text(slide, 0.9, 0.92, 11.6, 0.6, title, 25, INK, True)
    line(slide, 0.9, 1.62, 11.53, RULE_STRONG, 1)


def footer(slide, idx, total):
    text(slide, 0.9, 6.95, 6, 0.2, "MATERIAL MIND ENGINE · WEEKLY 2026-W42", 8, MUTED, name=MONO, spc=120)
    text(slide, 10.4, 6.95, 2.03, 0.2, f"{idx:02d} / {total:02d}", 8, MUTED, name=MONO,
         align=PP_ALIGN.RIGHT, spc=120)


def hair_table(slide, x, y, w, headers, rows, ratios, row_h=0.44, size=11, head_size=9.2):
    total = sum(ratios)
    xs = []
    cur = x
    for r in ratios:
        xs.append(cur)
        cur += w * r / total

    for i, h in enumerate(headers):
        cw = w * ratios[i] / total
        text(slide, xs[i], y, cw - 0.1, 0.22, h, head_size, MUTED, name=MONO, spc=100)
    line(slide, x, y + 0.28, w, RULE_STRONG, 1)

    top = y + 0.28
    for row in rows:
        for i, cell in enumerate(row):
            cw = w * ratios[i] / total
            color = INK if i == 0 else INK2
            text(slide, xs[i], top + 0.1, cw - 0.12, row_h - 0.1, str(cell), size, color,
                 bold=(i == 0), spacing=1.15)
        line(slide, x, top + row_h, w, RULE, 0.75)
        top += row_h
    return top


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    total = 11

    # 01 封面
    s = new_slide(prs)
    text(s, 0.9, 1.55, 11, 0.25, "WEEKLY REPORT · 2026-W42", 10, MUTED, name=MONO, spc=220)
    text(s, 0.9, 1.95, 11.5, 1.9, "Material Mind Engine", 46, INK, True, spacing=1.1)
    text(s, 0.9, 3.15, 10.5, 0.5, "材料数据自动化 RAG 引擎 · 第 42 周周报", 17, INK2)
    line(s, 0.9, 3.95, 11.53, RULE_STRONG, 1)
    text(s, 0.9, 4.15, 11, 0.3,
         "2026-10-18      v2.0.0      AGPL-3.0 + 商业双授权", 10.5, MUTED, name=MONO, spc=100)
    text(s, 0.9, 5.6, 11, 0.4, "从脚手架原型，迁移为可交付的私有化产品底座。", 14, INK2)
    footer(s, 1, total)

    # 02 概览
    s = new_slide(prs)
    header(s, "OVERVIEW", "本周概览")
    kpis = [
        ("代码变更", "+6,853", "−6,441 行 / 68 文件"),
        ("Python", "1,806", "core / ingest / server"),
        ("前端", "2,319", "零构建依赖"),
        ("已入库", "15", "条记录 / 3 个知识库"),
    ]
    for i, (label, value, note) in enumerate(kpis):
        x = 0.9 + i * 2.93
        text(s, x, 2.1, 2.7, 0.22, label, 9, MUTED, name=MONO, spc=140)
        text(s, x, 2.38, 2.7, 0.62, value, 32, INK, True, name=MONO)
        text(s, x, 3.02, 2.7, 0.24, note, 10, MUTED)
        if i < 3:
            vline(s, x + 2.78, 2.1, 1.15)
    line(s, 0.9, 3.62, 11.53, RULE_STRONG, 1)
    block(s, 0.9, 3.95, 11.4, [
        ("5 次提交 · 11 个 HTTP 接口 · 67 个受版本控制文件", "一次性完成架构替换、自动化投喂、资料库管理与商业授权落地。"),
        ("关键结论", "技术上已经能跑通全链路，但还证明不了产品价值——知识库只有 15 条记录。"),
    ])
    footer(s, 2, total)

    # 03 主线
    s = new_slide(prs)
    header(s, "WHAT CHANGED", "本周主线：一次架构级替换")
    text(s, 0.9, 1.95, 5.2, 0.25, "迁移前", 9.5, MUTED, name=MONO, spc=140)
    block(s, 0.9, 2.3, 5.2, [
        ("Chainlit 默认界面", "框架自带样式，能力上限是一个聊天框"),
        ("三套前端并存", "Chainlit / Gradio / Streamlit 各自维护"),
        ("入库需三步手动执行", "extract → clean → vector_store"),
        ("无进度、无通知", "失败只能翻日志"),
    ], gap=0.42)
    vline(s, 6.45, 1.95, 3.9)
    text(s, 6.9, 1.95, 5.4, 0.25, "迁移后", 9.5, ACCENT, name=MONO, spc=140)
    block(s, 6.9, 2.3, 5.5, [
        ("FastAPI + 自研前端", "视觉与交互完全可控，零构建依赖"),
        ("一套界面", "对话 + 资料库，同一个设计系统"),
        ("投喂即入库", "拖入文件自动完成六步流水线"),
        ("全程可见", "分阶段进度条 + 结果 toast + 失败原因"),
    ], gap=0.42)
    footer(s, 3, total)

    # 04 交付一
    s = new_slide(prs)
    header(s, "DELIVERED · 1 / 2", "交付：前端与投喂")
    block(s, 0.9, 1.95, 5.3, [
        ("前端重构", "顶部导航 + 集合标签页 + 资料库视图；技术文档式设计语言（纸张底色 / 衬线标题 / 细线分栏 / 单一强调色）。"),
        ("视觉系统", "18 个内联 SVG 线性图标；等宽数字；明暗双主题。"),
        ("动效", "视图切换、消息进场、表格 stagger、按压反馈，全部 160–240ms，只动 transform 与 opacity，支持系统减弱动效。"),
    ], gap=0.5)
    vline(s, 6.45, 1.95, 4.4)
    block(s, 6.9, 1.95, 5.5, [
        ("投喂流水线", "PDF / DOCX / XLSX / HTML / MD / CSV 解析 → 800 字分块 → LLM 11 字段结构化抽取 → 牌号与供应商归一化 → 内容哈希去重 → 向量化入库。"),
        ("inbox 守护", "目录监听，按子目录自动路由到对应知识库；成功归档 _done，失败归 _failed。"),
        ("低置信度隔离", "confidence < 0.6 不入主库，写入 quarantine.jsonl 等待人工确认。"),
    ], gap=0.5)
    footer(s, 4, total)

    # 05 交付二
    s = new_slide(prs)
    header(s, "DELIVERED · 2 / 2", "交付：管理与工程")
    block(s, 0.9, 1.95, 5.3, [
        ("资料库管理", "按文件聚合列表、文件名筛选、批量选择删除、JSON 导出、单文件重新解析、检索测试。"),
        ("引用可追溯", "答案内 [n] 角标可点击，定位并展开对应来源的 11 字段表与原始片段。"),
    ], gap=0.55)
    vline(s, 6.45, 1.95, 4.4)
    block(s, 6.9, 1.95, 5.5, [
        ("交互反馈", "全局入库队列、逐阶段进度条、+n 条 / 跳过 n 结果、成功与失败 toast、0 条时明确原因。"),
        ("工程与法务", "MIT → AGPL-3.0 + 商业双授权；新增 COMMERCIAL.md 与 CLA.md；重写 git 历史清除数据文件。"),
        ("发布", "tag v2.0.0；20 个 GitHub topics；仓库描述与主页信息补全。"),
    ], gap=0.5)
    footer(s, 5, total)

    # 06 指标
    s = new_slide(prs)
    header(s, "METRICS", "关键指标：迁移前后")
    hair_table(
        s, 0.9, 1.95, 11.5,
        ["指标", "迁移前", "当前", "变化"],
        [
            ["入库操作", "3 个脚本手动", "拖入即入库", "—"],
            ["答案长度", "284 字", "195 字", "−31%"],
            ["检索条数 top_k", "8", "5", "−37%"],
            ["分块大小", "1200 字", "800 字", "−33%"],
            ["引用可点击", "否", "是", "—"],
            ["入库可见进度", "无", "分阶段", "—"],
        ],
        [2.6, 2.6, 2.6, 1.6],
        row_h=0.52, size=12,
    )
    text(s, 0.9, 5.55, 11.4, 0.3,
         "答案改为紧凑列表后，输出 token 下降约 31%，同时更易扫读。", 11, MUTED)
    footer(s, 6, total)

    # 07 数据
    s = new_slide(prs)
    header(s, "CORPUS", "知识库现状：15 条")
    for i, (name, full, count, state) in enumerate([
        ("materials", "FFKM 全氟醚橡胶", "13", "待扩充"),
        ("peek", "PEEK 聚醚醚酮", "2", "待扩充"),
        ("fkm", "FKM 氟橡胶", "0", "空"),
    ]):
        y = 1.95 + i * 1.0
        text(s, 0.9, y, 2.4, 0.3, name, 12, MUTED, name=MONO, spc=80)
        text(s, 3.3, y - 0.1, 4.4, 0.4, full, 14, INK, True)
        text(s, 8.4, y - 0.12, 1.4, 0.4, count, 22, INK, True, name=MONO, align=PP_ALIGN.RIGHT)
        text(s, 10.0, y, 2.0, 0.3, state, 10, MUTED, name=MONO, align=PP_ALIGN.RIGHT, spc=80)
        line(s, 0.9, y + 0.62, 11.53, RULE, 0.75)
    text(s, 0.9, 5.15, 11.4, 0.8,
         "15 条记录足以验证链路，但不足以验证产品价值。数据量是当前最大短板，也是下周第一优先。",
         13, INK2)
    footer(s, 7, total)

    # 08 问题
    s = new_slide(prs)
    header(s, "FIXED", "问题与修复：6 个")
    hair_table(
        s, 0.9, 1.95, 11.5,
        ["问题", "根因"],
        [
            ["引用角标点击无反应", "两个脚本存在同名函数，来源卡片从未渲染"],
            ["投喂无任何反馈", "进度卡片被插入隐藏容器"],
            ["入库显示 +0 无原因", "后端返回了原因，前端未展示"],
            ["Embedding 调用 401", "环境变量不在配置白名单"],
            ["旧库查询静默降级", "缺少 embedding 模型元数据"],
            ["git 全部操作失败", "残留失效代理 + 缺证书链"],
        ],
        [2.4, 5.0],
        row_h=0.52, size=11.5,
    )
    footer(s, 8, total)

    # 09 风险
    s = new_slide(prs)
    header(s, "RISKS", "四个风险，按严重度排序")
    risks = [
        ("P0", "合规", "厂商 datasheet 的版权与数据源条款。只存事实性数值、保留来源链接、不复制长段落、遵守 robots.txt。这是唯一可能让项目归零的风险。"),
        ("P0", "数据量", "知识库只有 15 条记录，客户一试即穿。扩充真实资料是第一优先。"),
        ("P1", "无评测", "缺少评测集，调整分块、提示词、embedding 后无法判断变好还是变坏。"),
        ("P2", "供应商依赖", "单一模型与单一 embedding 服务商，接口变更会直接中断入库能力。"),
    ]
    y = 1.95
    for level, name, desc in risks:
        color = ACCENT if level == "P0" else MUTED
        text(s, 0.9, y + 0.02, 0.7, 0.25, level, 9.5, color, name=MONO, spc=100)
        text(s, 1.75, y, 1.9, 0.3, name, 13.5, INK, True)
        text(s, 3.75, y - 0.02, 8.7, 0.62, desc, 11.5, INK2, spacing=1.3)
        line(s, 0.9, y + 0.82, 11.53, RULE, 0.75)
        y += 1.12
    footer(s, 9, total)

    # 10 下周
    s = new_slide(prs)
    header(s, "NEXT WEEK · P0", "下周计划")
    items = [
        ("扩充真实资料", "FFKM / PEEK 各补 2–3 份厂商 datasheet，目标每个库 ≥ 50 条记录。入口：inbox/ 或网页拖入。"),
        ("建立评测集", "20 条「问题 → 期望数值 → 期望来源」，覆盖三种材料，作为后续所有改动的回归门禁。"),
        ("跑出基线数字", "输出检索命中率与数值准确率，形成一句可对外说的话：「20 题答对 16 题」。"),
        ("收敛双轨代码", "旧脚本与新 ingest/ 并存，标记 legacy 或删除，降低维护成本。"),
    ]
    y = 1.95
    for i, (title, desc) in enumerate(items, start=1):
        text(s, 0.9, y + 0.02, 0.55, 0.25, f"0{i}", 10, ACCENT, name=MONO, spc=120)
        text(s, 1.6, y - 0.02, 3.0, 0.3, title, 13.5, INK, True)
        text(s, 4.8, y - 0.02, 7.6, 0.62, desc, 11.5, INK2, spacing=1.3)
        line(s, 0.9, y + 0.8, 11.53, RULE, 0.75)
        y += 1.1
    text(s, 0.9, 6.32, 11.4, 0.35, "出口标准：20 题里 ≥ 16 题数值正确且引用对得上。", 13, INK, True)
    footer(s, 10, total)

    # 11 结尾
    s = new_slide(prs)
    text(s, 0.9, 2.5, 11, 0.25, "THANK YOU", 10, MUTED, name=MONO, spc=240)
    text(s, 0.9, 2.9, 11.5, 0.8, "Material Mind Engine", 38, INK, True, spacing=1.1)
    line(s, 0.9, 3.95, 11.53, RULE_STRONG, 1)
    text(s, 0.9, 4.2, 11, 0.3, "github.com/Yourstandby/Material-Mind-Engine", 12, INK2, name=MONO)
    text(s, 0.9, 4.55, 11, 0.3, "商业授权：luowei971024 [at] gmail.com", 12, MUTED)
    footer(s, 11, total)

    prs.save(OUT)
    print(f"saved: {OUT}  slides: {len(prs.slides.__iter__.__self__._sldIdLst)}")


if __name__ == "__main__":
    build()
